from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List

import pandas as pd
import streamlit as st

try:
    from docx import Document
except Exception:
    Document = None

DEFAULT_WORKFLOW: List[Dict[str, Any]] = [
    {"id":"prep","module":"稽查前准备","step_name":"确认资料齐全性与范围","instruction":"先确认资料是否齐全，并明确本次稽查重点。","ai_prompt":"请核对方案、ICF、伦理批件、授权表、培训记录、EDC导出表、既往问题清单。","documents":["方案","ICF与伦理批件","授权表","培训记录","EDC导出表","既往问题"],"fields":[{"key":"scope","label":"本次稽查范围/重点模块","type":"textarea"},{"key":"missing_docs","label":"缺失资料（如无填写：无）","type":"textarea"}],"common_risks":["范围不清导致漏查","资料缺失导致证据链不足"],"rules":["missing_docs"]},
    {"id":"icf","module":"知情同意","step_name":"核查版本、签署时序、授权与培训","instruction":"选择受试者，核对ICF版本、伦理批准、签署日期、首次研究操作、人员授权与培训。","ai_prompt":"现在查看ICF签署页，并与伦理批件、筛选检查、授权表、培训记录交叉核对。","documents":["ICF签署页","伦理批件","筛选检查记录","授权表","培训记录"],"fields":[{"key":"icf_version","label":"ICF版本号","type":"text"},{"key":"irb_approval_date","label":"伦理批准日期","type":"date"},{"key":"icf_sign_date","label":"受试者签署日期","type":"date"},{"key":"first_study_date","label":"首次研究相关操作日期","type":"date"},{"key":"signer","label":"签署/解释人员","type":"text"},{"key":"authorized","label":"是否已授权（是/否/不清楚）","type":"text"},{"key":"training_date","label":"相关培训日期","type":"date"}],"common_risks":["研究操作早于ICF签署","使用未获批版本","签署人员未授权","培训晚于操作"],"rules":["icf_timing","approval_timing","authorization","training_timing"]},
    {"id":"eligibility","module":"入排标准","step_name":"核查判断依据","instruction":"根据方案入排标准，确认入组判断前是否已有充分证据。","ai_prompt":"查看筛选检查、诊断记录、实验室、影像、既往用药及EDC入排页。","documents":["方案入排标准","筛选检查报告","病历诊断","实验室报告","EDC入排页"],"fields":[{"key":"random_date","label":"入组/随机日期","type":"date"},{"key":"latest_required_result_date","label":"最晚一项关键入排依据形成日期","type":"date"},{"key":"criteria_issue","label":"证据不足或异常说明","type":"textarea"}],"common_risks":["关键检查结果晚于入组判断","排除标准未充分排查","EDC与源文件不一致"],"rules":["eligibility_timing","text_issue"]},
    {"id":"visit","module":"访视与方案依从性","step_name":"核查访视窗口与关键操作","instruction":"核对方案窗口、实际访视日期、关键检查和漏做原因。","ai_prompt":"逐个访视确认日期是否在窗口内，关键检查是否完成。","documents":["方案访视表","源文件","EDC访视页","偏离记录"],"fields":[{"key":"visit_name","label":"访视名称","type":"text"},{"key":"window_requirement","label":"方案窗口要求","type":"text"},{"key":"actual_visit_date","label":"实际访视日期","type":"date"},{"key":"visit_issue","label":"超窗/漏做/顺序错误说明","type":"textarea"}],"common_risks":["访视超窗","关键检查漏做","EDC与源文件访视日期不一致"],"rules":["text_issue"]},
    {"id":"edc","module":"EDC与源文件一致性","step_name":"核查关键字段一致性","instruction":"抽查关键字段，确认EDC、源文件和Query记录是否一致。","ai_prompt":"选择一个关键字段进行核查。不一致时追溯源数据、query处理和研究者审核。","documents":["EDC导出表","源文件","Query记录","审计追踪"],"fields":[{"key":"edc_field","label":"核查字段/表单","type":"text"},{"key":"source_value","label":"源文件值","type":"text"},{"key":"edc_value","label":"EDC值","type":"text"},{"key":"query_status","label":"Query状态/处理说明","type":"textarea"}],"common_risks":["EDC与源文件不一致","query未查明根因","关键字段不一致"],"rules":["edc_mismatch"]},
    {"id":"summary","module":"问题汇总与报告初稿","step_name":"汇总问题和离场沟通","instruction":"复核全部问题记录，区分事实、影响、建议和待确认事项。","ai_prompt":"离场前确认每个问题是否有证据来源、是否需要进一步确认。","documents":["问题清单","证据截图","口头沟通纪要"],"fields":[{"key":"summary_note","label":"离场沟通重点/待确认事项","type":"textarea"}],"common_risks":["描述事实不足","建议措施不可执行","CAPA责任不清"],"rules":[]}
]

RISK_ORDER = {"低": 1, "中": 2, "高": 3}


def init_session() -> None:
    st.session_state.setdefault("current_index", 0)
    st.session_state.setdefault("done_steps", set())
    st.session_state.setdefault("step_inputs", {})
    st.session_state.setdefault("evidence_files", {})
    st.session_state.setdefault("records", [])
    st.session_state.setdefault("project", {})
    st.session_state.setdefault("last_result", None)


def get_current_step() -> Dict[str, Any]:
    return DEFAULT_WORKFLOW[st.session_state.current_index]


def next_step() -> None:
    st.session_state.current_index = min(st.session_state.current_index + 1, len(DEFAULT_WORKFLOW) - 1)
    st.session_state.last_result = None


def previous_step() -> None:
    st.session_state.current_index = max(st.session_state.current_index - 1, 0)
    st.session_state.last_result = None


def mark_step_done() -> None:
    st.session_state.done_steps.add(st.session_state.current_index)


def parse_date(value: str):
    if not value:
        return None
    value = value.strip().replace("/", "-").replace(".", "-")
    for fmt in ("%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    return None


def _add(res, risk, summary, question, rule):
    if RISK_ORDER[risk] > RISK_ORDER[res["risk_level"]]:
        res["risk_level"] = risk
    res["summaries"].append(summary)
    res["questions"].append(question)
    res["triggered_rules"].append(rule)
    res["need_finding"] = True


def evaluate_current_step() -> Dict[str, Any]:
    step = get_current_step()
    data = st.session_state.step_inputs.get(step["id"], {})
    res = {"risk_level": "低", "summaries": [], "questions": [], "triggered_rules": [], "need_finding": False}

    if "missing_docs" in step["rules"] and data.get("missing_docs", "").strip() not in ("", "无", "没有", "NA", "N/A"):
        _add(res, "中", "存在关键资料缺失，后续判断可能受限。", "请确认缺失资料是否可现场补充，并记录影响范围。", "资料缺失规则")

    if "text_issue" in step["rules"]:
        text = " ".join(str(v) for v in data.values()).strip()
        if text and not any(x in text for x in ["无", "未见异常", "正常", "无问题"]):
            _add(res, "中", "已记录潜在异常或证据不足事项。", "请补充受试者、文件名称、日期、版本或页码。", "文本异常提示规则")

    if "icf_timing" in step["rules"]:
        icf_date = parse_date(data.get("icf_sign_date", ""))
        first_date = parse_date(data.get("first_study_date", ""))
        if icf_date and first_date and first_date < icf_date:
            _add(res, "高", "首次研究相关操作日期早于ICF签署日期。", "请确认该操作是否为常规诊疗，并保留对应证据。", "ICF签署时序规则")

    if "approval_timing" in step["rules"]:
        approval = parse_date(data.get("irb_approval_date", ""))
        icf_date = parse_date(data.get("icf_sign_date", ""))
        if approval and icf_date and icf_date < approval:
            _add(res, "高", "受试者签署日期早于伦理批准日期。", "请核对ICF版本、伦理批件和签署页版本。", "ICF批准时序规则")

    if "authorization" in step["rules"] and data.get("authorized", "").strip() in ("否", "未授权", "不清楚", "未知"):
        _add(res, "中", "签署/解释人员授权状态不充分或不清楚。", "请查看授权表中该人员职责和生效日期。", "人员授权规则")

    if "training_timing" in step["rules"]:
        training = parse_date(data.get("training_date", ""))
        icf_date = parse_date(data.get("icf_sign_date", ""))
        if training and icf_date and training > icf_date:
            _add(res, "中", "培训日期晚于ICF签署日期。", "请确认是否存在更早培训记录。", "培训时序规则")

    if "eligibility_timing" in step["rules"]:
        random_date = parse_date(data.get("random_date", ""))
        latest_result = parse_date(data.get("latest_required_result_date", ""))
        if random_date and latest_result and latest_result > random_date:
            _add(res, "高", "关键入排判断依据形成日期晚于入组/随机日期。", "请确认研究者入组前是否已获得该结果。", "入排依据时序规则")

    if "edc_mismatch" in step["rules"]:
        source = str(data.get("source_value", "")).strip()
        edc = str(data.get("edc_value", "")).strip()
        if source and edc and source != edc:
            _add(res, "中", "EDC值与源文件值不一致。", "请补充字段、访视、源文件位置和Query状态。", "EDC一致性规则")

    if not res["summaries"]:
        res["summaries"].append("当前步骤未触发明确异常。")

    return {"step_id": step["id"], "module": step["module"], "risk_level": res["risk_level"], "summary": " ".join(res["summaries"]), "questions": res["questions"], "triggered_rules": res["triggered_rules"], "need_finding": res["need_finding"]}


def build_finding_from_current_step(result: Dict[str, Any]) -> Dict[str, str]:
    step = get_current_step()
    project = st.session_state.project
    data = st.session_state.step_inputs.get(step["id"], {})
    evidence_files = st.session_state.evidence_files.get(step["id"], [])
    evidence = list(step["documents"])
    if evidence_files:
        evidence.append("上传证据：" + "；".join(evidence_files))
    details = "；".join([f"{k}={v}" for k, v in data.items() if str(v).strip()])
    risk_impact = {"高":"可能影响受试者权益、入组判断、关键数据可靠性或核查解释。","中":"可能影响流程合规性、证据链完整性或数据一致性。","低":"当前风险较低，建议保留核查痕迹。"}
    rec = {"高":"建议立即补充证据并升级确认，必要时横向排查并制定CAPA。","中":"建议补充页码、截图和日期证据，明确责任人和完成期限。","低":"建议记录核查过程，后续结合更多样本判断。"}
    return {"项目名称":project.get("project_name", ""),"申办方":project.get("sponsor", ""),"中心编号":project.get("site_no", ""),"受试者编号":project.get("subject_id", ""),"稽查员":project.get("auditor", ""),"问题分类":step["module"],"问题标题":f"{step['module']}相关问题需确认","问题描述":result["summary"] + (" 现场记录信息：" + details if details else ""),"证据来源":"；".join(evidence),"风险等级":result["risk_level"],"风险影响":risk_impact[result["risk_level"]],"建议措施":rec[result["risk_level"]],"触发规则":"；".join(result["triggered_rules"]),"状态":"待稽查员确认","记录时间":datetime.now().strftime("%Y-%m-%d %H:%M")}


def export_records_dataframe() -> pd.DataFrame:
    return pd.DataFrame(st.session_state.records) if st.session_state.records else pd.DataFrame()


def build_word_report() -> bytes:
    if Document is None:
        return "未安装python-docx".encode("utf-8")
    doc = Document()
    p = st.session_state.project
    doc.add_heading("AI引导式稽查报告初稿", level=0)
    for k, label in [("project_name","项目名称"),("sponsor","申办方"),("site_no","中心编号"),("audit_type","稽查类型"),("auditor","稽查员")]:
        doc.add_paragraph(f"{label}：{p.get(k, '')}")
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_heading("一、稽查执行概况", level=1)
    doc.add_paragraph(f"系统共配置{len(DEFAULT_WORKFLOW)}个流程节点，已完成{len(st.session_state.done_steps)}个节点。")
    doc.add_heading("二、发现问题汇总", level=1)
    if not st.session_state.records:
        doc.add_paragraph("当前暂无已生成的问题记录。")
    for i, r in enumerate(st.session_state.records, 1):
        doc.add_heading(f"{i}. {r['问题分类']}｜{r['风险等级']}｜{r['问题标题']}", level=2)
        for key in ["问题描述","证据来源","风险影响","建议措施","状态"]:
            doc.add_paragraph(f"{key}：{r.get(key, '')}")
    doc.add_heading("三、离场沟通与CAPA建议", level=1)
    doc.add_paragraph(st.session_state.step_inputs.get("summary", {}).get("summary_note", "建议形成离场沟通清单，明确责任人、完成期限和复核方式。"))
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
